#!/usr/bin/env python3
"""Fig S8B: motif-annotated TF-footprint density over the IDO1 TSS +/-100 bp window.

Combines sequence motif hits, FiberHMM TF-sized protected calls, condition-level footprint
gains and RNA-seq expression support, and writes the ranked candidate tables.

Also the motif-scan engine that make_FigS8D.py imports for the CXCL11 element window.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pysam
from scipy.ndimage import gaussian_filter1d
from scipy.stats import fisher_exact

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared"))

import config as cfg
import figstyle as fs

FIBERHMM_DIR = cfg.path("dafseq", "fiberhmm_dir")
SAMPLE_SETS = Path(__file__).resolve().parents[1] / "data" / "locus_sample_sets.tsv"

CHROM = "chr8"
TSS = 39_913_890
WINDOW_START = TSS - 100
WINDOW_END = TSS + 100

PANEL_ID = "FigS8B"

SAMPLES = OrderedDict(
    [
        ("Naive", "GB-270_IDO1_Naive_merged"),
        ("Stimulated", "GB-270_IDO1_Stim_merged"),
        ("Rested", "GB-270_IDO1_Rested_merged"),
        ("Rechallenged", "GB-270_IDO1_Rechallenge_merged"),
    ]
)
LOCUS_LABEL = "IDO1"
OUTPUT_PREFIX = "ido1_tss_pm100"

RNA_CONDITION_MAP = OrderedDict(
    [
        ("Naive", "NTC"),
        ("Stimulated", "IFNg"),
        ("Rested", "Trained"),
        ("Rechallenged", "Rechallenge"),
    ]
)

COLORS = {
    "Naive": "#909499",
    "Stimulated": "#4D9AD4",
    "Rested": "#4A4A4B",
    "Rechallenged": "#0B519D",
}

FAMILY_COLORS = {
    "STAT/ISGF3": "#C44E52",
    "IRF": "#8172B3",
    "AP-1": "#DD8452",
    "NF-kB/REL": "#55A868",
    "CEBP": "#CCB974",
    "SP/KLF/EGR": "#4C72B0",
    "Other": "#8C8C8C",
}

BG_DEFAULT = np.array([0.25, 0.25, 0.25, 0.25], dtype=float)
BASE_IDX = {"A": 0, "C": 1, "G": 2, "T": 3}
COMPLEMENT = str.maketrans("ACGTacgt", "TGCAtgca")
PSEUDO_TPM = 0.1


@dataclass
class Motif:
    motif_id: str
    name: str
    width: int
    nsites: int
    pwm: np.ndarray
    max_score: float
    min_score: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Rank candidate TF identities for FiberHMM protected calls in a locus-specific window."
    )
    parser.add_argument("--locus", default="IDO1", help="Locus label used for sample config, titles, and output names")
    parser.add_argument("--region", default=f"{CHROM}:{WINDOW_START}-{WINDOW_END}")
    parser.add_argument("--tss", type=int, default=TSS)
    parser.add_argument("--sample-set", type=Path, default=SAMPLE_SETS)
    parser.add_argument("--window-label", default="tss_pm100")
    parser.add_argument("--reference", type=Path, default=cfg.path("reference", "hg38_fasta"))
    parser.add_argument("--motifs", type=Path, default=cfg.path("reference", "jaspar_motifs"))
    parser.add_argument("--tpm-xlsx", type=Path, default=cfg.path("rnaseq", "tpm_xlsx"))
    parser.add_argument("--counts-xlsx", type=Path, default=cfg.path("rnaseq", "counts"))
    parser.add_argument("--pwm-threshold", type=float, default=0.85)
    parser.add_argument("--out-dir", type=Path, default=None)
    parser.add_argument("--top-motifs", type=int, default=10)
    parser.add_argument(
        "--extra-plot-tf",
        action="append",
        default=[],
        help="TF name to force-add to motif tracks with --extra-plot-threshold; may be repeated.",
    )
    parser.add_argument(
        "--extra-plot-threshold",
        type=float,
        default=0.75,
        help="PWM score fraction threshold for --extra-plot-tf annotations.",
    )
    return parser.parse_args()


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def load_locus_samples(locus: str, sample_set: Path) -> OrderedDict[str, str]:
    rows = [row for row in read_tsv(sample_set) if row["locus"].upper() == locus.upper()]
    rows.sort(key=lambda row: int(row["order"]))
    if not rows:
        raise ValueError(f"No samples configured for locus {locus!r} in {sample_set}")
    return OrderedDict((row["condition"], row["sample"]) for row in rows)


def parse_region(region: str) -> tuple[str, int, int]:
    chrom, coords = region.replace(",", "").split(":")
    start, end = coords.split("-")
    return chrom, int(start), int(end)


def normalize_gene_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9-]", "", name).upper()


def motif_components(tf_name: str) -> list[str]:
    cleaned = tf_name.strip()
    if "::" in cleaned:
        raw_parts = cleaned.split("::")
    elif "/" in cleaned:
        raw_parts = cleaned.split("/")
    else:
        raw_parts = [cleaned]
    parts = []
    for part in raw_parts:
        gene = normalize_gene_name(part)
        if gene:
            parts.append(gene)
    return parts or [normalize_gene_name(tf_name)]


def motif_family(tf_name: str) -> str:
    name = normalize_gene_name(tf_name)
    comps = motif_components(tf_name)
    if any(c.startswith("STAT") for c in comps) or "ISGF3" in name:
        return "STAT/ISGF3"
    if any(c.startswith("IRF") for c in comps):
        return "IRF"
    if any(c in {"RELA", "RELB", "REL", "NFKB1", "NFKB2"} or c.startswith("NFKB") for c in comps):
        return "NF-kB/REL"
    if any(c in {"JUN", "JUNB", "JUND", "FOS", "FOSB", "FOSL1", "FOSL2", "BATF"} for c in comps):
        return "AP-1"
    if any(c.startswith("CEBP") for c in comps):
        return "CEBP"
    if any(c.startswith("SP") or c.startswith("KLF") or c.startswith("EGR") or c == "MAZ" for c in comps):
        return "SP/KLF/EGR"
    return "Other"


def expression_rank(tier: str) -> int:
    return {
        "high_ifng_induced": 4,
        "expressed_induced": 3,
        "expressed_not_induced": 2,
        "low_expression": 1,
        "missing_expression": 0,
    }.get(tier, 0)


def reverse_complement(seq: str) -> str:
    return seq.translate(COMPLEMENT)[::-1]


def parse_background(lines: list[str]) -> np.ndarray:
    for i, line in enumerate(lines):
        if line.startswith("Background letter frequencies") and i + 1 < len(lines):
            parts = lines[i + 1].split()
            values = {}
            for base, value in zip(parts[::2], parts[1::2]):
                values[base.upper()] = float(value)
            if all(base in values for base in "ACGT"):
                return np.array([values["A"], values["C"], values["G"], values["T"]], dtype=float)
    return BG_DEFAULT.copy()


def parse_meme_motifs(meme_path: Path) -> list[Motif]:
    content = meme_path.read_text()
    lines = content.splitlines()
    background = parse_background(lines)
    motifs = []
    for block in re.split(r"\nMOTIF ", content)[1:]:
        block_lines = block.strip().splitlines()
        if not block_lines:
            continue
        header = block_lines[0].strip().split()
        motif_id = header[0]
        name = header[1] if len(header) > 1 else motif_id

        mat_start = None
        width = 0
        nsites = 20
        for i, line in enumerate(block_lines):
            if line.strip().startswith("letter-probability matrix"):
                width_match = re.search(r"w=\s*(\d+)", line)
                nsites_match = re.search(r"nsites=\s*(\d+)", line)
                if width_match:
                    width = int(width_match.group(1))
                if nsites_match:
                    nsites = int(nsites_match.group(1))
                mat_start = i + 1
                break
        if mat_start is None or width <= 0:
            continue

        ppm_rows = []
        for line in block_lines[mat_start:]:
            stripped = line.strip()
            if not stripped or stripped.startswith(("URL", "MOTIF", "letter", "Background")):
                break
            vals = stripped.split()
            if len(vals) != 4:
                break
            try:
                ppm_rows.append([float(v) for v in vals])
            except ValueError:
                break
        if len(ppm_rows) != width:
            continue

        ppm = np.array(ppm_rows, dtype=float)
        ppm = (ppm * nsites + 0.25) / (nsites + 1.0)
        with np.errstate(divide="ignore"):
            pwm = np.log2(ppm / background)
        pwm = np.where(np.isfinite(pwm), pwm, -20.0)
        motifs.append(
            Motif(
                motif_id=motif_id,
                name=name,
                width=width,
                nsites=nsites,
                pwm=pwm,
                max_score=float(pwm.max(axis=1).sum()),
                min_score=float(pwm.min(axis=1).sum()),
            )
        )
    return motifs


def encode_seq(seq: str) -> np.ndarray:
    arr = np.full(len(seq), -1, dtype=np.int8)
    for i, base in enumerate(seq.upper()):
        arr[i] = BASE_IDX.get(base, -1)
    return arr


def score_position(pwm: np.ndarray, seq_arr: np.ndarray, pos: int) -> float:
    score = 0.0
    for i in range(pwm.shape[0]):
        idx = seq_arr[pos + i]
        if idx >= 0:
            score += pwm[i, idx]
    return float(score)


def scan_motifs(motifs: list[Motif], seq: str, region_start: int, threshold_frac: float) -> list[dict]:
    seq = seq.upper()
    seq_rc = reverse_complement(seq)
    seq_enc = encode_seq(seq)
    seq_rc_enc = encode_seq(seq_rc)
    region_len = len(seq)
    hits = []

    for motif in motifs:
        width = motif.width
        if region_len < width:
            continue
        span = motif.max_score - motif.min_score
        threshold = motif.min_score + threshold_frac * span if span > 0 else motif.max_score * 0.95
        for pos in range(region_len - width + 1):
            score = score_position(motif.pwm, seq_enc, pos)
            if score >= threshold:
                score_frac = (score - motif.min_score) / span if span > 0 else 1.0
                hits.append(
                    {
                        "motif_id": motif.motif_id,
                        "tf_name": motif.name,
                        "abs_start": region_start + pos,
                        "abs_end": region_start + pos + width,
                        "motif_strand": "+",
                        "score": round(score, 6),
                        "score_frac": round(score_frac, 6),
                        "width": width,
                    }
                )
            score = score_position(motif.pwm, seq_rc_enc, pos)
            if score >= threshold:
                fwd_pos = region_len - pos - width
                score_frac = (score - motif.min_score) / span if span > 0 else 1.0
                hits.append(
                    {
                        "motif_id": motif.motif_id,
                        "tf_name": motif.name,
                        "abs_start": region_start + fwd_pos,
                        "abs_end": region_start + fwd_pos + width,
                        "motif_strand": "-",
                        "score": round(score, 6),
                        "score_frac": round(score_frac, 6),
                        "width": width,
                    }
                )
    for i, hit in enumerate(sorted(hits, key=lambda h: (h["abs_start"], h["abs_end"], h["tf_name"], h["motif_strand"]))):
        hit["motif_uid"] = f"motif_{i + 1:04d}"
        hit["family"] = motif_family(hit["tf_name"])
        hit["components"] = ",".join(motif_components(hit["tf_name"]))
    return sorted(hits, key=lambda h: h["motif_uid"])


def sample_paths(sample: str) -> dict[str, Path]:
    return {
        "recalled_bam": FIBERHMM_DIR / f"{sample}.fiberhmm.recalled.bam",
        "tf_bed": FIBERHMM_DIR / f"{sample}.fiberhmm.recalled_tf.bed",
    }


def parse_bed12_blocks(path: Path):
    with path.open() as handle:
        for line in handle:
            if not line.strip() or line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            chrom = fields[0]
            start = int(fields[1])
            read = fields[3]
            score = int(fields[4])
            strand = fields[5]
            if len(fields) >= 12:
                block_count = int(fields[9])
                sizes = [int(x) for x in fields[10].rstrip(",").split(",") if x]
                starts = [int(x) for x in fields[11].rstrip(",").split(",") if x]
                for size, rel in zip(sizes[:block_count], starts[:block_count]):
                    block_start = start + rel
                    yield {
                        "chrom": chrom,
                        "start": block_start,
                        "end": block_start + size,
                        "read": read,
                        "score": score,
                        "strand": strand,
                        "size": size,
                    }
            else:
                end = int(fields[2])
                yield {
                    "chrom": chrom,
                    "start": start,
                    "end": end,
                    "read": read,
                    "score": score,
                    "strand": strand,
                    "size": end - start,
                }


def load_blocks(path: Path, chrom: str, start: int, end: int) -> list[dict]:
    return [
        block
        for block in parse_bed12_blocks(path)
        if block["chrom"] == chrom and block["end"] > start and block["start"] < end
    ]


def eligible_reads(bam_path: Path, chrom: str, start: int, end: int) -> set[str]:
    reads = set()
    with pysam.AlignmentFile(str(bam_path), "rb") as bam:
        for read in bam.fetch(chrom, start, end):
            if read.is_unmapped or read.is_secondary or read.is_supplementary:
                continue
            if read.reference_start <= start and read.reference_end >= end:
                reads.add(read.query_name)
    return reads


def load_condition_data(chrom: str, start: int, end: int) -> dict[str, dict]:
    data = {}
    for condition, sample in SAMPLES.items():
        paths = sample_paths(sample)
        eligible = eligible_reads(paths["recalled_bam"], chrom, start, end)
        blocks = load_blocks(paths["tf_bed"], chrom, start, end)
        blocks_by_read = defaultdict(list)
        for block in blocks:
            if block["read"] in eligible:
                blocks_by_read[block["read"]].append(block)
        data[condition] = {
            "sample": sample,
            "eligible": eligible,
            "blocks": blocks,
            "blocks_by_read": blocks_by_read,
            "n_eligible": len(eligible),
        }
    return data


def fraction(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else float("nan")


def safe_log2_ratio(num: float, den: float, pseudo: float = PSEUDO_TPM) -> float:
    return math.log2((num + pseudo) / (den + pseudo))


def bh_qvalues(p_values: list[float]) -> list[float]:
    n = len(p_values)
    indexed = sorted(enumerate(p_values), key=lambda item: item[1])
    q_values = [float("nan")] * n
    prev = 1.0
    for rank, (idx, p_value) in enumerate(reversed(indexed), start=1):
        true_rank = n - rank + 1
        q = min(prev, p_value * n / true_rank)
        q_values[idx] = q
        prev = q
    return q_values


def tsv_float(value: float, digits: int = 6) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "NA"
    return f"{value:.{digits}f}"


def load_tpm_expression(path: Path) -> dict[str, dict[str, float]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["all data"]
    rows = ws.iter_rows(values_only=True)
    header1 = next(rows)
    header2 = next(rows)
    avg_cols = {}
    for idx, (cond, metric) in enumerate(zip(header1, header2)):
        cond = str(cond).strip() if cond is not None else ""
        metric = str(metric).strip() if metric is not None else ""
        if metric == "AVG TPM" and cond in {"NTC", "IFNg", "Trained", "Rechallenge"}:
            avg_cols[cond] = idx
    missing = {"NTC", "IFNg", "Trained", "Rechallenge"} - set(avg_cols)
    if missing:
        raise ValueError(f"Missing expected TPM columns in {path}: {sorted(missing)}")

    expression = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        gene = normalize_gene_name(str(row[0]))
        if not gene:
            continue
        expression[gene] = {
            cond: float(row[col] or 0.0)
            for cond, col in avg_cols.items()
        }
    return expression


def load_count_expression(path: Path) -> dict[str, dict[str, float]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    cond_cols = defaultdict(list)
    for idx, value in enumerate(header):
        label = str(value).strip() if value is not None else ""
        for cond in ("NTC", "IFNg", "Trained", "Rechallenge"):
            if cond in label:
                cond_cols[cond].append(idx)
    missing = {"NTC", "IFNg", "Trained", "Rechallenge"} - set(cond_cols)
    if missing:
        raise ValueError(f"Missing expected count columns in {path}: {sorted(missing)}")

    counts = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        gene = normalize_gene_name(str(row[0]))
        if not gene:
            continue
        counts[gene] = {}
        for cond, cols in cond_cols.items():
            vals = [float(row[col] or 0.0) for col in cols]
            counts[gene][cond] = float(np.mean(vals)) if vals else 0.0
    return counts


def expression_annotation(tf_name: str, tpm: dict[str, dict[str, float]], counts: dict[str, dict[str, float]]) -> dict:
    components = motif_components(tf_name)
    component_tpm = []
    component_counts = []
    missing = []
    for comp in components:
        if comp not in tpm:
            missing.append(comp)
        else:
            component_tpm.append(tpm[comp])
            component_counts.append(counts.get(comp, {}))

    if missing or not component_tpm:
        return {
            "components": ",".join(components),
            "expression_components_found": ",".join([c for c in components if c not in missing]),
            "missing_expression_components": ",".join(missing),
            "rna_naive_tpm": float("nan"),
            "rna_stimulated_tpm": float("nan"),
            "rna_rested_tpm": float("nan"),
            "rna_rechallenged_tpm": float("nan"),
            "rna_ifng_vs_ntc_log2fc": float("nan"),
            "rna_rechallenge_vs_trained_log2fc": float("nan"),
            "rna_max_tpm": float("nan"),
            "raw_count_present": False,
            "raw_count_max_avg": float("nan"),
            "expression_tier": "missing_expression",
        }

    tpm_by_daf = {}
    for daf_cond, rna_cond in RNA_CONDITION_MAP.items():
        vals = [entry.get(rna_cond, 0.0) for entry in component_tpm]
        tpm_by_daf[daf_cond] = min(vals) if len(vals) > 1 else vals[0]

    count_max_by_component = []
    for comp_counts in component_counts:
        count_max_by_component.append(max(comp_counts.values()) if comp_counts else 0.0)
    raw_count_max_avg = min(count_max_by_component) if len(count_max_by_component) > 1 else count_max_by_component[0]
    raw_count_present = raw_count_max_avg >= 10

    ifng_log2fc = safe_log2_ratio(tpm_by_daf["Stimulated"], tpm_by_daf["Naive"])
    rechallenge_log2fc = safe_log2_ratio(tpm_by_daf["Rechallenged"], tpm_by_daf["Rested"])
    max_tpm = max(tpm_by_daf.values())

    if tpm_by_daf["Stimulated"] >= 50 and ifng_log2fc >= 1:
        tier = "high_ifng_induced"
    elif tpm_by_daf["Stimulated"] >= 5 and ifng_log2fc >= 0.5:
        tier = "expressed_induced"
    elif max_tpm >= 5:
        tier = "expressed_not_induced"
    else:
        tier = "low_expression"

    return {
        "components": ",".join(components),
        "expression_components_found": ",".join(components),
        "missing_expression_components": "",
        "rna_naive_tpm": tpm_by_daf["Naive"],
        "rna_stimulated_tpm": tpm_by_daf["Stimulated"],
        "rna_rested_tpm": tpm_by_daf["Rested"],
        "rna_rechallenged_tpm": tpm_by_daf["Rechallenged"],
        "rna_ifng_vs_ntc_log2fc": ifng_log2fc,
        "rna_rechallenge_vs_trained_log2fc": rechallenge_log2fc,
        "rna_max_tpm": max_tpm,
        "raw_count_present": raw_count_present,
        "raw_count_max_avg": raw_count_max_avg,
        "expression_tier": tier,
    }


def annotate_hits_with_expression(hits: list[dict], tpm: dict, counts: dict) -> dict[str, dict]:
    annotations = {}
    for hit in hits:
        if hit["tf_name"] not in annotations:
            annotations[hit["tf_name"]] = expression_annotation(hit["tf_name"], tpm, counts)
    return annotations


def compute_occupancy(hits: list[dict], condition_data: dict[str, dict]) -> tuple[list[dict], list[dict]]:
    occupancy_rows = []
    summary_by_hit = {hit["motif_uid"]: dict(hit) for hit in hits}
    p_stim = []
    p_rechallenge = []

    for hit in hits:
        center = (hit["abs_start"] + hit["abs_end"]) // 2
        per_condition = {}
        for condition, data in condition_data.items():
            center_reads = set()
            interval_reads = set()
            for read_name, blocks in data["blocks_by_read"].items():
                for block in blocks:
                    if block["start"] <= center < block["end"]:
                        center_reads.add(read_name)
                    if block["end"] > hit["abs_start"] and block["start"] < hit["abs_end"]:
                        interval_reads.add(read_name)
            n_eligible = data["n_eligible"]
            n_center = len(center_reads)
            n_interval = len(interval_reads)
            per_condition[condition] = {
                "eligible_fibers": n_eligible,
                "center_overlap_fibers": n_center,
                "interval_overlap_fibers": n_interval,
                "center_overlap_fraction": fraction(n_center, n_eligible),
                "interval_overlap_fraction": fraction(n_interval, n_eligible),
            }
            occupancy_rows.append(
                {
                    **hit,
                    "condition": condition,
                    "sample": data["sample"],
                    "eligible_fibers": n_eligible,
                    "center_overlap_fibers": n_center,
                    "interval_overlap_fibers": n_interval,
                    "center_overlap_fraction": fraction(n_center, n_eligible),
                    "interval_overlap_fraction": fraction(n_interval, n_eligible),
                    "relative_start": hit["abs_start"] - TSS,
                    "relative_end": hit["abs_end"] - TSS,
                    "relative_center": center - TSS,
                }
            )

        naive = per_condition["Naive"]
        stim = per_condition["Stimulated"]
        rested = per_condition["Rested"]
        rechallenged = per_condition["Rechallenged"]

        stim_table = [
            [stim["center_overlap_fibers"], stim["eligible_fibers"] - stim["center_overlap_fibers"]],
            [naive["center_overlap_fibers"], naive["eligible_fibers"] - naive["center_overlap_fibers"]],
        ]
        rechallenge_table = [
            [
                rechallenged["center_overlap_fibers"],
                rechallenged["eligible_fibers"] - rechallenged["center_overlap_fibers"],
            ],
            [rested["center_overlap_fibers"], rested["eligible_fibers"] - rested["center_overlap_fibers"]],
        ]
        stim_odds, stim_p = fisher_exact(stim_table, alternative="greater")
        rechallenge_odds, rechallenge_p = fisher_exact(rechallenge_table, alternative="greater")
        p_stim.append(float(stim_p))
        p_rechallenge.append(float(rechallenge_p))

        summary = summary_by_hit[hit["motif_uid"]]
        for condition, metrics in per_condition.items():
            prefix = condition.lower()
            summary[f"{prefix}_eligible_fibers"] = metrics["eligible_fibers"]
            summary[f"{prefix}_center_overlap_fibers"] = metrics["center_overlap_fibers"]
            summary[f"{prefix}_interval_overlap_fibers"] = metrics["interval_overlap_fibers"]
            summary[f"{prefix}_center_overlap_fraction"] = metrics["center_overlap_fraction"]
            summary[f"{prefix}_interval_overlap_fraction"] = metrics["interval_overlap_fraction"]
        summary["stim_vs_naive_delta"] = (
            stim["center_overlap_fraction"] - naive["center_overlap_fraction"]
        )
        summary["rechallenge_vs_rested_delta"] = (
            rechallenged["center_overlap_fraction"] - rested["center_overlap_fraction"]
        )
        summary["stim_vs_naive_odds_ratio"] = float(stim_odds)
        summary["stim_vs_naive_p"] = float(stim_p)
        summary["rechallenge_vs_rested_odds_ratio"] = float(rechallenge_odds)
        summary["rechallenge_vs_rested_p"] = float(rechallenge_p)
        summary["relative_start"] = hit["abs_start"] - TSS
        summary["relative_end"] = hit["abs_end"] - TSS
        summary["relative_center"] = center - TSS

    q_stim = bh_qvalues(p_stim)
    q_rechallenge = bh_qvalues(p_rechallenge)
    summary_rows = list(summary_by_hit.values())
    for row, q1, q2 in zip(summary_rows, q_stim, q_rechallenge):
        row["stim_vs_naive_q"] = q1
        row["rechallenge_vs_rested_q"] = q2
    return occupancy_rows, summary_rows


def add_candidate_scores(summary_rows: list[dict], expr_by_tf: dict[str, dict]) -> list[dict]:
    scored = []
    for row in summary_rows:
        expr = expr_by_tf.get(row["tf_name"], {})
        tier = expr.get("expression_tier", "missing_expression")
        family_bonus = {
            "STAT/ISGF3": 0.7,
            "IRF": 0.6,
            "NF-kB/REL": 0.4,
            "AP-1": 0.35,
            "CEBP": 0.25,
            "SP/KLF/EGR": 0.15,
            "Other": 0.0,
        }.get(row["family"], 0.0)
        score = (
            12.0 * max(row["stim_vs_naive_delta"], 0.0)
            + 7.0 * max(row["rechallenge_vs_rested_delta"], 0.0)
            + 0.8 * float(row["score_frac"])
            + 0.3 * expression_rank(tier)
            + family_bonus
        )
        out = {
            **row,
            **expr,
            "priority_score": score,
            "support_summary": support_summary(row, expr),
        }
        scored.append(out)

    scored.sort(
        key=lambda r: (
            -r["priority_score"],
            -r["stim_vs_naive_delta"],
            -r["stimulated_center_overlap_fraction"],
            -float(r["score_frac"]),
        )
    )
    for rank, row in enumerate(scored, start=1):
        row["rank"] = rank
    return scored


def support_summary(row: dict, expr: dict) -> str:
    parts = []
    if row["stim_vs_naive_delta"] > 0.02:
        parts.append("stim_gain")
    if row["rechallenge_vs_rested_delta"] > 0.02:
        parts.append("memory_reactivation_gain")
    if row["family"] in {"STAT/ISGF3", "IRF", "NF-kB/REL", "AP-1", "CEBP"}:
        parts.append("ifn_msc_plausible_family")
    if expr.get("expression_tier") in {"high_ifng_induced", "expressed_induced"}:
        parts.append(expr["expression_tier"])
    elif expr.get("expression_tier") == "expressed_not_induced":
        parts.append("expressed")
    if not parts:
        parts.append("sequence_footprint_candidate")
    return ";".join(parts)


def collapse_families(candidate_rows: list[dict]) -> list[dict]:
    grouped = defaultdict(list)
    for row in candidate_rows:
        grouped[row["family"]].append(row)

    family_rows = []
    for family, rows in grouped.items():
        top = min(rows, key=lambda r: r["rank"])
        family_rows.append(
            {
                "family": family,
                "n_motif_hits": len(rows),
                "top_rank": top["rank"],
                "top_tf_name": top["tf_name"],
                "top_components": top["components"],
                "top_relative_start": top["relative_start"],
                "top_relative_end": top["relative_end"],
                "top_motif_score_frac": top["score_frac"],
                "best_priority_score": top["priority_score"],
                "max_stimulated_fraction": max(r["stimulated_center_overlap_fraction"] for r in rows),
                "max_stim_vs_naive_delta": max(r["stim_vs_naive_delta"] for r in rows),
                "max_rechallenge_vs_rested_delta": max(r["rechallenge_vs_rested_delta"] for r in rows),
                "best_expression_tier": max(rows, key=lambda r: expression_rank(r["expression_tier"]))[
                    "expression_tier"
                ],
                "representative_support": top["support_summary"],
            }
        )
    family_rows.sort(key=lambda r: (-r["best_priority_score"], r["family"]))
    return family_rows


def save_tsv(path: Path, rows: list[dict], fieldnames: list[str] | None = None) -> None:
    if not rows:
        raise ValueError(f"No rows to write for {path}")
    if fieldnames is None:
        seen = []
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.append(key)
        fieldnames = seen
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            formatted = {}
            for key in fieldnames:
                value = row.get(key, "")
                if isinstance(value, float):
                    formatted[key] = tsv_float(value)
                else:
                    formatted[key] = value
            writer.writerow(formatted)


def expression_rows(expr_by_tf: dict[str, dict], hits: list[dict]) -> list[dict]:
    rows = []
    seen = set()
    for hit in sorted(hits, key=lambda h: (h["family"], h["tf_name"])):
        tf_name = hit["tf_name"]
        if tf_name in seen:
            continue
        seen.add(tf_name)
        rows.append(
            {
                "tf_name": tf_name,
                "family": hit["family"],
                **expr_by_tf[tf_name],
            }
        )
    return rows


def save_figure(fig: plt.Figure, out_dir: Path, stem: str) -> None:
    fs.save(fig, out_dir / stem)


def extra_plot_motif_hits(
    motifs: list[Motif],
    seq: str,
    region_start: int,
    tss: int,
    tf_names: list[str],
    threshold_frac: float,
) -> list[dict]:
    """Scan requested TF motifs for plot-only annotations without changing rankings."""
    targets = {normalize_gene_name(name) for name in tf_names if normalize_gene_name(name)}
    if not targets:
        return []
    extra_motifs = [
        motif
        for motif in motifs
        if normalize_gene_name(motif.name) in targets
        or any(component in targets for component in motif_components(motif.name))
    ]
    hits = scan_motifs(extra_motifs, seq, region_start, threshold_frac)
    for i, hit in enumerate(hits, start=1):
        hit["motif_uid"] = f"extra_plot_{i:04d}"
        hit["relative_start"] = hit["abs_start"] - tss
        hit["relative_end"] = hit["abs_end"] - tss
        hit["relative_center"] = ((hit["abs_start"] + hit["abs_end"]) // 2) - tss
        hit["extra_plot_motif"] = True
        hit["plot_label"] = f"{hit['tf_name']} ({hit['score_frac']:.2f})"
    return hits


def select_plot_motifs(
    candidate_rows: list[dict],
    top_n: int,
    gap: int = 6,
    extra_rows: list[dict] | None = None,
) -> list[dict]:
    """Choose readable representative motif labels without hiding full TSV detail."""
    selected = []
    for row in candidate_rows:
        overlaps_existing = any(
            row["relative_start"] < kept["relative_end"] + gap
            and row["relative_end"] > kept["relative_start"] - gap
            for kept in selected
        )
        if overlaps_existing:
            continue
        selected.append(row)
        if len(selected) >= top_n:
            break
    if extra_rows:
        for row in extra_rows:
            already_selected = any(
                row["motif_id"] == kept["motif_id"]
                and row["relative_start"] == kept["relative_start"]
                and row["relative_end"] == kept["relative_end"]
                and row["motif_strand"] == kept["motif_strand"]
                for kept in selected
            )
            if not already_selected:
                selected.append(row)
    return selected


def per_base_density(condition_data: dict, start: int, end: int) -> dict[str, np.ndarray]:
    width = end - start
    densities = {}
    for condition, data in condition_data.items():
        counts = np.zeros(width, dtype=float)
        for read_name, blocks in data["blocks_by_read"].items():
            for block in blocks:
                s = max(block["start"], start) - start
                e = min(block["end"], end) - start
                if s < e:
                    counts[s:e] += 1
        densities[condition] = counts / max(data["n_eligible"], 1)
    return densities


def draw_top_motif_track(ax, top_rows: list[dict], start: int, tss: int) -> None:
    ax.set_ylim(0, 1)
    ax.set_yticks([])
    ax.spines[["left", "right", "top"]].set_visible(False)
    ax.tick_params(axis="x", labelsize=6)
    n_levels = min(max(len(top_rows), 1), 5)
    y_levels = np.linspace(0.13, 0.72, n_levels)
    for i, row in enumerate(top_rows):
        y = y_levels[i % len(y_levels)] if len(y_levels) else 0.5
        x0 = row["abs_start"] - tss
        x1 = row["abs_end"] - tss
        color = FAMILY_COLORS.get(row["family"], FAMILY_COLORS["Other"])
        is_extra = row.get("extra_plot_motif", False)
        ax.add_patch(
            plt.Rectangle(
                (x0, y - 0.06),
                x1 - x0,
                0.12,
                facecolor=color,
                edgecolor="#222222" if is_extra else color,
                alpha=0.55 if is_extra else 0.85,
                linewidth=0.5 if is_extra else 0,
            )
        )
        label = row.get("plot_label")
        if not label:
            label = row["tf_name"] if row["family"] == "Other" else f"{row['family']} {row['tf_name']}"
        ax.text(
            (x0 + x1) / 2,
            y + 0.09,
            label,
            fontsize=6,
            ha="center",
            va="bottom",
            rotation=30,
            color="#222222",
        )


def plot_motif_density(
    condition_data: dict,
    candidate_rows: list[dict],
    out_dir: Path,
    start: int,
    end: int,
    tss: int,
    top_n: int,
    extra_plot_rows: list[dict] | None = None,
) -> None:
    density = per_base_density(condition_data, start, end)
    plot_rows = select_plot_motifs(candidate_rows, top_n, extra_rows=extra_plot_rows)
    x = np.arange(start, end) - tss
    fig, (ax_track, ax) = plt.subplots(
        2,
        1,
        figsize=(8.2, 5.1),
        sharex=True,
        gridspec_kw={"height_ratios": [1.2, 3.0], "hspace": 0.04},
    )
    draw_top_motif_track(ax_track, plot_rows, start, tss)
    ax_track.set_ylabel("Top\nmotifs", fontsize=8, rotation=0, labelpad=24, va="center")
    for condition in SAMPLES:
        y = gaussian_filter1d(density[condition], sigma=2)
        ax.plot(x, y, color=COLORS[condition], lw=2.0, label=condition)
    ax.axvline(0, color="#222222", lw=1.0, ls="--")
    ax.axvspan(-75, 75, color="#777777", alpha=0.07, lw=0)
    ax.set_xlim(start - tss, end - tss)
    ax.set_ylim(bottom=0)
    ax.set_xlabel(f"Position relative to {LOCUS_LABEL} TSS (bp)")
    ax.set_ylabel("Fraction of spanning fibers\nwith FiberHMM TF call")
    ax.legend(frameon=False, ncols=4, fontsize=6)
    fig.suptitle(f"{LOCUS_LABEL} TF window: FiberHMM protection density with top candidate motifs", fontsize=8, y=1.04)
    save_figure(fig, out_dir, PANEL_ID)


def motif_hit_fields() -> list[str]:
    return [
        "motif_uid",
        "motif_id",
        "tf_name",
        "family",
        "components",
        "abs_start",
        "abs_end",
        "relative_start",
        "relative_end",
        "relative_center",
        "motif_strand",
        "score",
        "score_frac",
        "width",
    ]


def candidate_fields() -> list[str]:
    return [
        "rank",
        "priority_score",
        "motif_uid",
        "motif_id",
        "tf_name",
        "family",
        "components",
        "abs_start",
        "abs_end",
        "relative_start",
        "relative_end",
        "relative_center",
        "motif_strand",
        "score_frac",
        "naive_eligible_fibers",
        "naive_center_overlap_fibers",
        "naive_center_overlap_fraction",
        "stimulated_eligible_fibers",
        "stimulated_center_overlap_fibers",
        "stimulated_center_overlap_fraction",
        "rested_eligible_fibers",
        "rested_center_overlap_fibers",
        "rested_center_overlap_fraction",
        "rechallenged_eligible_fibers",
        "rechallenged_center_overlap_fibers",
        "rechallenged_center_overlap_fraction",
        "stim_vs_naive_delta",
        "stim_vs_naive_odds_ratio",
        "stim_vs_naive_p",
        "stim_vs_naive_q",
        "rechallenge_vs_rested_delta",
        "rechallenge_vs_rested_odds_ratio",
        "rechallenge_vs_rested_p",
        "rechallenge_vs_rested_q",
        "rna_naive_tpm",
        "rna_stimulated_tpm",
        "rna_rested_tpm",
        "rna_rechallenged_tpm",
        "rna_ifng_vs_ntc_log2fc",
        "rna_rechallenge_vs_trained_log2fc",
        "rna_max_tpm",
        "raw_count_present",
        "raw_count_max_avg",
        "expression_tier",
        "missing_expression_components",
        "support_summary",
    ]


def main() -> None:
    args = parse_args()
    global LOCUS_LABEL, OUTPUT_PREFIX, SAMPLES, TSS
    LOCUS_LABEL = args.locus.upper()
    OUTPUT_PREFIX = f"{LOCUS_LABEL.lower()}_{args.window_label}"
    SAMPLES = load_locus_samples(LOCUS_LABEL, args.sample_set)
    TSS = args.tss
    chrom, start, end = parse_region(args.region)
    out_dir = args.out_dir or cfg.out_dir(PANEL_ID)
    out_dir.mkdir(parents=True, exist_ok=True)
    fs.setup()

    with pysam.FastaFile(str(args.reference)) as fasta:
        seq = fasta.fetch(chrom, start, end)
    motifs = parse_meme_motifs(args.motifs)
    hits = scan_motifs(motifs, seq, start, args.pwm_threshold)
    for hit in hits:
        hit["relative_start"] = hit["abs_start"] - args.tss
        hit["relative_end"] = hit["abs_end"] - args.tss
        hit["relative_center"] = ((hit["abs_start"] + hit["abs_end"]) // 2) - args.tss

    tpm = load_tpm_expression(args.tpm_xlsx)
    counts = load_count_expression(args.counts_xlsx)
    expr_by_tf = annotate_hits_with_expression(hits, tpm, counts)

    condition_data = load_condition_data(chrom, start, end)
    occupancy_rows, summary_rows = compute_occupancy(hits, condition_data)
    candidate_rows = add_candidate_scores(summary_rows, expr_by_tf)
    family_rows = collapse_families(candidate_rows)
    extra_plot_rows = extra_plot_motif_hits(
        motifs,
        seq,
        start,
        args.tss,
        args.extra_plot_tf,
        args.extra_plot_threshold,
    )

    save_tsv(out_dir / f"{OUTPUT_PREFIX}_motif_hits.tsv", hits, motif_hit_fields())
    save_tsv(out_dir / f"{OUTPUT_PREFIX}_motif_occupancy_by_condition.tsv", occupancy_rows)
    save_tsv(out_dir / f"{OUTPUT_PREFIX}_candidate_tf_ranking.tsv", candidate_rows, candidate_fields())
    save_tsv(out_dir / f"{OUTPUT_PREFIX}_candidate_families.tsv", family_rows)
    save_tsv(out_dir / f"{OUTPUT_PREFIX}_expression_annotation.tsv", expression_rows(expr_by_tf, hits))

    plot_motif_density(
        condition_data,
        candidate_rows,
        out_dir,
        start,
        end,
        args.tss,
        args.top_motifs,
        extra_plot_rows,
    )

    print(f"Region: {chrom}:{start}-{end} ({start - args.tss}..{end - args.tss} bp relative to TSS)")
    print(f"Motifs scanned: {len(motifs):,}")
    print(f"Motif hits: {len(hits):,}")
    for condition, data in condition_data.items():
        print(f"{condition}: {data['n_eligible']:,} fully spanning fibers")
    print("Top candidates:")
    for row in candidate_rows[:10]:
        print(
            f"  {row['rank']:>2}. {row['family']:<10} {row['tf_name']:<18} "
            f"{row['relative_start']:>4}..{row['relative_end']:<4} "
            f"stim_delta={row['stim_vs_naive_delta']:.4f} "
            f"expr={row['expression_tier']}"
        )
    if args.extra_plot_tf:
        print("Extra plot motif hits:")
        if extra_plot_rows:
            for row in extra_plot_rows:
                print(
                    f"  {row['family']:<10} {row['tf_name']:<18} "
                    f"{row['relative_start']:>4}..{row['relative_end']:<4} "
                    f"strand={row['motif_strand']} score_frac={row['score_frac']:.4f}"
                )
        else:
            print("  none")
    print(f"Wrote outputs to {out_dir}")


if __name__ == "__main__":
    main()
